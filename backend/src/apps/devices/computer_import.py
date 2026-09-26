"""
Kompyuterlarni Excel'dan ommaviy qo'shish.

Ustunlar (nomdagi qavs ichi e'tiborsiz qoldiriladi -
`dtm_id (region.dtm_id)` = `dtm_id`):

    dtm_id          viloyat (`Region.dtm_id`)          majburiy
    zone_number     bino (`Zone.number`, o'sha viloyatda) majburiy
    machine_uuid    mashina (SMBIOS UUID) - ASOSIY      majburiy
    mac_address     tarmoq kartasi                     ixtiyoriy
    number          xonadagi tartib raqami             majburiy
    inventory_code  buxgalteriya kodi                  ixtiyoriy

Viloyat va bino TASHQI RAQAMLAR bilan beriladi (FaceID integratsiyasidagi
kabi, `CLAUDE.md`): ichki ID'ni Excel tayyorlaydigan odam bilmaydi.

HAMMASI YOKI HECH NARSA. Bitta qatorda xato bo'lsa, birorta ham qator
yozilmaydi va javobda har bir xato qator/ustun bilan qaytadi. Qisman
import "qaysilari tushdi?" degan savolni qoldirardi va tuzatilgan faylni
qayta yuklash endi yarmi "allaqachon bor" bo'lib chiqardi. Shu sababli
`dry_run` ham bor: panel avval tekshiradi, keyin yozadi - ikkalasi bitta
funksiya, ya'ni tekshiruvda "o'tdi" degan fayl yozishda yiqilmaydi
(poyga holatidan tashqari, uni bazadagi cheklovlar ushlaydi).

UUID ALLAQACHON RO'YXATDA - XATO EMAS, O'TKAZIB YUBORILADI. Tuzatilgan
yoki to'ldirilgan faylni qayta yuklash odatiy ish va u oldingi safar
qo'shilgan mashinalar tufayli yiqilmasligi kerak. Mavjud yozuv
O'ZGARTIRILMAYDI: import - qo'shish vositasi, tahrirlash emas.

YAGONA ISTISNO - UUID'NI TO'LDIRISH. UUID'dan oldingi (MAC bilan
qo'shilgan) kompyuter qatordagi MAC bo'yicha topilsa va uning UUID'i
BO'SH bo'lsa, faqat `machine_uuid` yoziladi (raqam, kod, bino
o'zgarmaydi). Ya'ni eski faylga bitta ustun qo'shib qayta yuklash butun
bino inventarini yangi identifikatorga o'tkazadi. Bino mos kelmasa -
xato: bu boshqa binoning mashinasi bo'lishi mumkin.
"""

from __future__ import annotations

import io
import re

from django.db import transaction
from django.utils import timezone

from apps.common.utils.validators import normalize_mac, normalize_machine_uuid
from apps.regions.models import Region, Zone

from .models import Computer
from .services import auto_inventory_code

MAX_ROWS = 5000
MAX_FILE_BYTES = 5 * 1024 * 1024

#: (ustun kodi, shablondagi sarlavha, majburiymi)
COLUMNS = (
    ("dtm_id", "dtm_id", True),
    ("zone_number", "zone_number", True),
    ("machine_uuid", "machine_uuid", True),
    ("mac_address", "mac_address", False),
    ("number", "number", True),
    ("inventory_code", "inventory_code", False),
)
_CODES = {code for code, _, _ in COLUMNS}
_REQUIRED = [code for code, _, required in COLUMNS if required]
_INVENTORY_RE = re.compile(r"^[A-Za-z0-9_\-]{3,50}$")


class ImportFileError(Exception):
    """Fayl umuman o'qilmadi (qator darajasidagi xato emas)."""


def _clean(value) -> str:
    """Katak qiymati -> satr; qavs ichi tashlanadi (`12 (Toshkent)` -> `12`)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel sonni float qilib saqlaydi: 12 -> 12.0.
        value = int(value)
    return str(value).split("(")[0].strip()


def _header(value) -> str:
    return _clean(value).lower().replace(" ", "_")


def _int(value: str):
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------
# Shablon
# --------------------------------------------------------------------------
def build_template() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Kompyuterlar"
    ws.append([title for _, title, _ in COLUMNS])
    head_fill = PatternFill("solid", fgColor="1F6F5C")
    optional_fill = PatternFill("solid", fgColor="6B7C77")
    for index, (_, _, required) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill if required else optional_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = 28
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    ws.column_dimensions["C"].width = 42  # UUID 36 belgi
    # UUID, MAC va inventar kodi MATN bo'lishi kerak: Excel "00-1A-..."
    # yoki "0012" ni son/sana deb, "1E10-..." ni esa ilmiy son deb
    # o'girib qo'yardi.
    for letter in ("C", "D", "F"):
        for row in range(2, MAX_ROWS + 2):
            ws[f"{letter}{row}"].number_format = "@"
    positive = DataValidation(type="whole", operator="greaterThan", formula1="0", allow_blank=True)
    positive.error = "Musbat butun son kiriting"
    ws.add_data_validation(positive)
    positive.add(f"A2:B{MAX_ROWS + 1}")
    positive.add(f"E2:E{MAX_ROWS + 1}")

    guide = wb.create_sheet("Yo'riqnoma")
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 80
    rows = [
        ("Ustun", "Qiymat"),
        ("dtm_id", "Viloyatning DTM ID raqami (Viloyatlar sahifasidagi «DTM ID»). Majburiy."),
        ("zone_number", "Binoning raqami - o'sha viloyat ichida (Binolar sahifasi). Majburiy."),
        ("machine_uuid", "Mashinaning ASOSIY identifikatori - ona platadagi SMBIOS UUID "
                         "(XXXXXXXX-XXXX-XXXX-XXXX-XXXXXXXXXXXX). Majburiy, tizim bo'ylab unikal. "
                         "Mashinada: `wmic csproduct get uuid` yoki PowerShell "
                         "`(Get-CimInstance Win32_ComputerSystemProduct).UUID`. Client o'rnatilgan "
                         "bo'lsa - panel «Qurilmalar» sahifasida (mashina UUID'i)."),
        ("mac_address", "Ixtiyoriy. AA:BB:CC:DD:EE:FF yoki AA-BB-CC-DD-EE-FF. Berilsa tizim bo'ylab "
                        "unikal. MAC o'zgarishi mumkin (tarmoq kartasi), shuning uchun u identifikator emas."),
        ("number", "Xonadagi tartib raqami (stoldagi raqam), 1..32767. Majburiy, bino ichida unikal."),
        ("inventory_code", "Ixtiyoriy. 3-50 belgi: lotin harfi, raqam, '-', '_'. "
                           "Bo'sh bo'lsa AUTO-<UUID> ko'rinishida yaratiladi."),
        ("", ""),
        ("Qoidalar", "Sarlavhadagi qavs ichi e'tiborsiz qoldiriladi. Bitta qatorda xato bo'lsa "
                     "hech narsa yozilmaydi. UUID allaqachon ro'yxatda bo'lsa qator o'tkazib "
                     "yuboriladi (mavjud yozuv o'zgarmaydi). Ko'pi bilan {} qator.".format(MAX_ROWS)),
        ("UUID'siz eski yozuvlar", "Kompyuter ilgari MAC bilan qo'shilgan va UUID'i bo'sh bo'lsa, "
                                   "shu MAC va o'sha bino ko'rsatilgan qator unga FAQAT UUID yozadi "
                                   "(raqam va kod o'zgarmaydi)."),
        ("Namuna", "dtm_id=10, zone_number=1, machine_uuid=4C4C4544-0038-4A10-805A-C7C04F4B3A12, "
                   "mac_address=00:1A:2B:3C:4D:5E, number=12, inventory_code=INV-0012"),
    ]
    for row in rows:
        guide.append(row)
    guide["A1"].font = Font(bold=True)
    guide["B1"].font = Font(bold=True)
    for row in guide.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].font = Font(bold=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# O'qish
# --------------------------------------------------------------------------
def read_rows(upload) -> list[dict]:
    """Birinchi varaqni o'qiydi: `[{"row": 2, "dtm_id": "10", ...}, ...]`."""
    from openpyxl import load_workbook

    if upload.size > MAX_FILE_BYTES:
        raise ImportFileError("Fayl juda katta (ko'pi bilan 5 MB).")
    try:
        wb = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Fayl o'qilmadi - .xlsx formatidagi Excel fayl yuklang.")
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ImportFileError("Fayl bo'sh.")
        index = {}
        for position, value in enumerate(header):
            code = _header(value)
            if code in _CODES and code not in index:
                index[code] = position
        missing = [code for code in _REQUIRED if code not in index]
        if missing:
            raise ImportFileError("Ustun topilmadi: {}.".format(", ".join(missing)))

        result = []
        for number, values in enumerate(rows, start=2):
            item = {
                code: _clean(values[position]) if position < len(values) else ""
                for code, position in index.items()
            }
            if not any(item.values()):
                continue  # bo'sh qator (Excel oxirida qoldiradi)
            item["row"] = number
            result.append(item)
            if len(result) > MAX_ROWS:
                raise ImportFileError("Qatorlar juda ko'p (ko'pi bilan {}).".format(MAX_ROWS))
        if not result:
            raise ImportFileError("Faylda ma'lumot qatori yo'q.")
        return result
    finally:
        wb.close()


# --------------------------------------------------------------------------
# Tekshirish va yozish
# --------------------------------------------------------------------------
def import_computers(rows: list[dict], *, user, dry_run: bool) -> dict:
    """
    Qatorlarni tekshiradi va (xato bo'lmasa, `dry_run=False` da) yozadi.

    Barcha bazaviy qidiruvlar TO'PLAM bo'yicha - qator soniga bog'liq
    bo'lmagan bir necha so'rov (viloyat, bino, UUID, MAC, inventar, raqam).
    """
    errors: list[dict] = []
    skipped: list[dict] = []

    def fail(row, column, message):
        errors.append({"row": row["row"], "column": column, "message": message})

    dtm_ids = {_int(r.get("dtm_id")) for r in rows} - {None}
    regions = {r.dtm_id: r for r in Region.objects.filter(dtm_id__in=dtm_ids)}
    zones = {
        (z.region_id, z.number): z
        for z in Zone.objects.filter(region__in=regions.values()).select_related("region")
    }

    parsed = []
    for row in rows:
        before = len(errors)
        dtm_id = _int(row.get("dtm_id"))
        region = regions.get(dtm_id)
        zone = None
        if dtm_id is None:
            fail(row, "dtm_id", "Viloyat DTM ID si kiritilmagan yoki son emas.")
        elif region is None:
            fail(row, "dtm_id", "DTM ID {} li viloyat topilmadi.".format(dtm_id))
        elif user.is_region_scoped and region.pk != user.region_id:
            fail(row, "dtm_id", "Bu viloyat sizning hududingizga tegishli emas.")
        else:
            zone_number = _int(row.get("zone_number"))
            zone = zones.get((region.pk, zone_number))
            if zone_number is None:
                fail(row, "zone_number", "Bino raqami kiritilmagan yoki son emas.")
            elif zone is None:
                fail(row, "zone_number", "«{}» viloyatida {}-bino topilmadi.".format(region.name, zone_number))

        machine_uuid = normalize_machine_uuid(row.get("machine_uuid"))
        if not row.get("machine_uuid"):
            fail(row, "machine_uuid", "Machine UUID kiritilmagan.")
        elif not machine_uuid:
            fail(row, "machine_uuid", "Machine UUID noto'g'ri yoki to'ldirilmagan "
                                      "(XXXXXXXX-XXXX-XXXX-XXXX-XXXXXXXXXXXX).")

        mac = None
        if row.get("mac_address"):
            mac = normalize_mac(row.get("mac_address"))
            if mac is None:
                fail(row, "mac_address", "MAC formati noto'g'ri (AA:BB:CC:DD:EE:FF).")

        number = _int(row.get("number"))
        if number is None or not 1 <= number <= 32767:
            fail(row, "number", "Kompyuter raqami 1..32767 oralig'idagi butun son bo'lishi kerak.")

        code = row.get("inventory_code", "")
        auto = not code
        if code and not _INVENTORY_RE.match(code):
            fail(row, "inventory_code", "Inventar kodi: 3-50 belgi, lotin harfi, raqam, '-' yoki '_'.")
        if not code and machine_uuid:
            code = auto_inventory_code(machine_uuid=machine_uuid)

        if len(errors) == before:
            parsed.append({"row": row["row"], "zone": zone, "uuid": machine_uuid, "mac": mac,
                           "number": number, "code": code, "auto": auto})

    # --- fayl ichidagi takrorlar -----------------------------------------
    def duplicates(key, column, message):
        seen = {}
        for item in parsed:
            value = key(item)
            if value is None:
                continue
            if value in seen:
                errors.append({"row": item["row"], "column": column,
                               "message": message.format(seen[value])})
            else:
                seen[value] = item["row"]

    duplicates(lambda i: i["uuid"], "machine_uuid", "Bu Machine UUID faylda takrorlangan ({}-qator).")
    duplicates(lambda i: i["mac"], "mac_address", "Bu MAC faylda takrorlangan ({}-qator).")
    # UUID dan hosil qilingan kod tekshirilmaydi: uning takrori - UUID
    # takrori va u yuqorida allaqachon xato bo'ldi (bitta sabab uchun
    # ikkita xato jadvalni chalkashtirardi).
    duplicates(lambda i: None if i["auto"] else i["code"].upper(), "inventory_code", "Bu inventar kodi faylda takrorlangan ({}-qator).")
    duplicates(lambda i: (i["zone"].pk, i["number"]), "number", "Bu raqam shu binoda faylda takrorlangan ({}-qator).")

    # --- bazadagi mavjud yozuvlar -----------------------------------------
    alive = Computer.objects.filter(deleted_at__isnull=True)
    existing_uuid = {
        c.machine_uuid: c
        for c in alive.filter(machine_uuid__in={i["uuid"] for i in parsed}).select_related("zone")
    }
    existing_mac = {
        c.mac_address: c
        for c in alive.filter(mac_address__in={i["mac"] for i in parsed if i["mac"]}).select_related("zone")
    }
    to_create, to_bind = [], []
    for item in parsed:
        known = existing_uuid.get(item["uuid"])
        if known is not None:
            skipped.append({"row": item["row"], "machine_uuid": item["uuid"],
                            "message": "Allaqachon ro'yxatda: {} ({}).".format(known.label, known.zone.name)})
            continue
        by_mac = existing_mac.get(item["mac"]) if item["mac"] else None
        if by_mac is None:
            to_create.append(item)
        elif by_mac.machine_uuid:
            errors.append({"row": item["row"], "column": "mac_address",
                           "message": "Bu MAC {} kompyuterida band (uning UUID'i boshqa: {}).".format(
                               by_mac.label, by_mac.machine_uuid)})
        elif by_mac.zone_id != item["zone"].pk:
            errors.append({"row": item["row"], "column": "zone_number",
                           "message": "Bu MAC «{}» binosidagi {} kompyuteriga tegishli.".format(
                               by_mac.zone.name, by_mac.label)})
        else:
            to_bind.append({**item, "computer": by_mac})

    taken_codes = set(
        alive.filter(inventory_code__in={i["code"] for i in to_create})
        .values_list("inventory_code", flat=True)
    )
    zone_ids = {i["zone"].pk for i in to_create}
    taken_numbers = set(
        alive.filter(zone_id__in=zone_ids, number__in={i["number"] for i in to_create})
        .values_list("zone_id", "number")
    )
    for item in to_create:
        if item["code"] in taken_codes:
            errors.append({"row": item["row"], "column": "inventory_code",
                           "message": "«{}» inventar kodi boshqa kompyuterda band.".format(item["code"])})
        if (item["zone"].pk, item["number"]) in taken_numbers:
            errors.append({"row": item["row"], "column": "number",
                           "message": "{}-raqam «{}» binosida band.".format(item["number"], item["zone"].name)})

    errors.sort(key=lambda e: (e["row"], e["column"]))
    report = {
        "total": len(rows),
        "to_create": 0 if errors else len(to_create),
        #: UUID'si bo'sh mavjud kompyuterlarga faqat UUID yoziladi.
        "to_bind": 0 if errors else len(to_bind),
        "binds": [
            {"row": i["row"], "machine_uuid": i["uuid"],
             "message": "UUID yoziladi: {} ({}).".format(i["computer"].label, i["computer"].zone.name)}
            for i in to_bind
        ],
        "skipped": skipped,
        "errors": errors,
        "created": 0,
        "bound": 0,
        "dry_run": dry_run,
    }
    if errors or dry_run or not (to_create or to_bind):
        return report

    with transaction.atomic():
        Computer.objects.bulk_create([
            Computer(zone=item["zone"], machine_uuid=item["uuid"], mac_address=item["mac"] or "",
                     number=item["number"], inventory_code=item["code"], is_active=True)
            for item in to_create
        ])
        for item in to_bind:
            # `machine_uuid IS NULL` sharti - tekshiruvdan keyin handshake
            # UUID'ni bog'lab ulgurgan bo'lsa, uning ustidan yozilmaydi.
            Computer.objects.filter(pk=item["computer"].pk, machine_uuid__isnull=True).update(
                machine_uuid=item["uuid"], updated_at=timezone.now()
            )
    report["created"] = len(to_create)
    report["bound"] = len(to_bind)
    return report
