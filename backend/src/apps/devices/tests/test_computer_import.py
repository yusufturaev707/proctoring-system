"""
Kompyuterlarni Excel'dan ommaviy qo'shish (`devices/computer_import.py`).

Mashina - (machine_uuid, mac_address) JUFTLIGI, ikkala ustun majburiy.
UUID bir partiyada takrorlanadi: UUID bir xil, MAC boshqa qator - yangi
kompyuter. MAC UUID'siz eski yozuvga UUID yozish uchun ham ishlatiladi.
"""

import io
import re
import zipfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.devices.models import Computer
from apps.proctoring.models import AuditLog
from apps.proctoring.tests import factories
from apps.proctoring.tests.test_client_api import bearer

HEADER = [
    "dtm_id (region.dtm_id)", "zone_number (zone.number)", "machine_uuid",
    "mac_address", "number", "inventory_code",
]
U1 = "4C4C4544-0038-4A10-805A-C7C04F4B3A01"
U2 = "4C4C4544-0038-4A10-805A-C7C04F4B3A02"
U3 = "4C4C4544-0038-4A10-805A-C7C04F4B3A03"
M1 = "00:1A:2B:3C:4D:01"
M2 = "00:1A:2B:3C:4D:02"
M3 = "00:1A:2B:3C:4D:03"


def xlsx(rows, header=HEADER) -> SimpleUploadedFile:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("k.xlsx", buffer.getvalue())


class ComputerImportTests(TestCase):
    def setUp(self):
        self.region = factories.make_region()
        self.zone = factories.make_zone(region=self.region, number=3)
        self.user = factories.make_user(permissions=["devices.view", "devices.manage"])
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=bearer(self.user))
        self.url = reverse("computer-import-excel")

    def post(self, rows, dry_run=False, **kwargs):
        return self.client.post(
            self.url, {"file": xlsx(rows, **kwargs), "dry_run": str(dry_run).lower()},
            format="multipart",
        )

    def test_creates_rows_in_region_zone(self):
        dtm = self.region.dtm_id
        response = self.post([
            [dtm, 3, U1.lower(), "00-1a-2b-3c-4d-5e", 12, "INV-0012"],
            [f"{dtm} (Viloyat)", "3 (Bino)", "{" + U2 + "}", M2.lower(), 13, None],
        ])
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["data"]["created"], 2)
        # UUID kanonik shaklda (katta harf, qavssiz) - client aynan shunday yuboradi.
        first = Computer.objects.get(machine_uuid=U1)
        self.assertEqual(
            (first.zone, first.number, first.inventory_code, first.mac_address),
            (self.zone, 12, "INV-0012", "00:1A:2B:3C:4D:5E"),
        )
        self.assertIsNone(first.ip_address)
        second = Computer.objects.get(number=13)
        self.assertEqual(second.machine_uuid, U2)
        self.assertEqual(second.mac_address, M2)
        # Inventar kodi UUID va MAC dan: faqat UUID'dan olingan kod bir
        # partiyadagi ikkinchi mashinada to'qnashardi.
        self.assertEqual(second.inventory_code, "AUTO-" + U2.replace("-", "") + "-" + M2.replace(":", ""))
        self.assertTrue(AuditLog.objects.filter(action="import").exists())

    def test_dry_run_writes_nothing(self):
        response = self.post([[self.region.dtm_id, 3, U1, M1, 1, ""]], dry_run=True)
        self.assertEqual(response.json()["data"]["to_create"], 1)
        self.assertFalse(Computer.objects.exists())

    def test_any_error_blocks_whole_file(self):
        response = self.post([
            [self.region.dtm_id, 3, U1, M1, 1, ""],
            [self.region.dtm_id, 99, U2, M2, 2, ""],                              # bino yo'q
            [self.region.dtm_id, 3, "yaroqsiz", M3, 3, ""],                       # UUID
            [self.region.dtm_id, 3, U2, M1, 4, ""],                               # takror MAC
            [self.region.dtm_id, 3, "00000000-0000-0000-0000-000000000000", "00:1A:2B:3C:4D:05", 5, ""],
            [self.region.dtm_id, 3, U3, "not-a-mac", 6, ""],                      # MAC formati
            [self.region.dtm_id, 3, U3, "", 7, ""],                               # MAC yo'q
        ])
        errors = response.json()["data"]["errors"]
        self.assertEqual({(e["row"], e["column"]) for e in errors}, {
            (3, "zone_number"), (4, "machine_uuid"), (5, "mac_address"),
            (6, "machine_uuid"), (7, "mac_address"), (8, "mac_address"),
        })
        self.assertFalse(Computer.objects.exists())

    def test_existing_pair_is_skipped_not_error(self):
        factories.make_computer(zone=self.zone, machine_uuid=U1, mac_address=M1, number=1)
        data = self.post([
            [self.region.dtm_id, 3, U1, M1.lower(), 1, ""],
            [self.region.dtm_id, 3, U2, M2, 2, ""],
        ]).json()["data"]
        self.assertEqual((data["created"], len(data["skipped"]), data["errors"]), (1, 1, []))

    def test_same_uuid_with_other_mac_is_a_new_computer(self):
        """Bir partiyadagi platalar - o'tkazib yuborilmaydi, yangi yozuv."""
        factories.make_computer(zone=self.zone, machine_uuid=U1, mac_address=M1, number=1)
        data = self.post([
            [self.region.dtm_id, 3, U1, M2, 2, ""],
            [self.region.dtm_id, 3, U1, M3, 3, ""],
        ]).json()["data"]
        self.assertEqual((data["created"], data["skipped"], data["errors"]), (2, [], []))
        self.assertEqual(
            set(Computer.objects.filter(machine_uuid=U1).values_list("mac_address", flat=True)),
            {M1, M2, M3},
        )

    def test_empty_mac_is_error(self):
        errors = self.post([[self.region.dtm_id, 3, U1, "", 1, ""]]).json()["data"]["errors"]
        self.assertEqual([(e["row"], e["column"]) for e in errors], [(2, "mac_address")])
        self.assertFalse(Computer.objects.exists())

    def test_legacy_record_without_mac_with_same_uuid_is_error(self):
        """
        MAC'siz eski yozuv - xato: u shu mashina bo'lishi mumkin, yangi
        yozuv esa uni UUID bo'yicha tanilmaydigan qilib qo'yardi.
        """
        legacy = factories.make_computer(zone=self.zone, machine_uuid=U1, mac_address="", number=1)
        errors = self.post([[self.region.dtm_id, 3, U1, M1, 2, ""]]).json()["data"]["errors"]
        self.assertEqual(errors[0]["column"], "mac_address")
        self.assertIn(legacy.label, errors[0]["message"])
        self.assertEqual(Computer.objects.count(), 1)

    def test_fills_uuid_of_legacy_computer_found_by_mac(self):
        """
        UUID'dan oldingi yozuv: MAC va bino mos - FAQAT UUID yoziladi.

        Raqam va kod o'zgarmaydi: import tahrirlash vositasi emas, eski
        faylga ustun qo'shib qayta yuklash bino inventarini UUID'ga
        o'tkazishi kerak.
        """
        legacy = factories.make_computer(
            zone=self.zone, machine_uuid=None, mac_address="00:1A:2B:3C:4D:5E", number=4,
        )
        dry = self.post([[self.region.dtm_id, 3, U1, "00-1A-2B-3C-4D-5E", 9, ""]], dry_run=True)
        self.assertEqual((dry.json()["data"]["to_bind"], dry.json()["data"]["to_create"]), (1, 0))

        data = self.post([[self.region.dtm_id, 3, U1, "00-1A-2B-3C-4D-5E", 9, ""]]).json()["data"]

        self.assertEqual((data["bound"], data["created"], data["errors"]), (1, 0, []))
        legacy.refresh_from_db()
        self.assertEqual((legacy.machine_uuid, legacy.number), (U1, 4))
        self.assertEqual(Computer.objects.count(), 1)

    def test_legacy_mac_in_other_building_is_error(self):
        other = factories.make_zone(region=self.region, number=4)
        factories.make_computer(zone=other, machine_uuid=None, mac_address="00:1A:2B:3C:4D:5E")
        errors = self.post([[self.region.dtm_id, 3, U1, "00:1A:2B:3C:4D:5E", 9, ""]]).json()["data"]["errors"]
        self.assertEqual(errors[0]["column"], "zone_number")

    def test_mac_of_computer_with_other_uuid_is_error(self):
        factories.make_computer(zone=self.zone, machine_uuid=U2, mac_address="00:1A:2B:3C:4D:5E")
        errors = self.post([[self.region.dtm_id, 3, U1, "00:1A:2B:3C:4D:5E", 9, ""]]).json()["data"]["errors"]
        self.assertEqual(errors[0]["column"], "mac_address")

    def test_taken_number_in_zone_is_error(self):
        factories.make_computer(zone=self.zone, number=7)
        errors = self.post([[self.region.dtm_id, 3, U1, M1, 7, ""]]).json()["data"]["errors"]
        self.assertEqual(errors[0]["column"], "number")

    def test_missing_required_column(self):
        """Eski shablon (UUID ustunisiz) - fayl darajasidagi xato."""
        response = self.post(
            [[1, 2, "00:1A:2B:3C:4D:5E", 3]],
            header=["dtm_id", "zone_number", "mac_address", "number"],
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("machine_uuid", response.content.decode())
        self.assertFalse(Computer.objects.exists())

    def test_region_admin_cannot_import_other_region(self):
        other = factories.make_region()
        factories.make_zone(region=other, number=1)
        user = factories.make_user(permissions=["devices.view", "devices.manage"], region=self.region)
        self.client.credentials(HTTP_AUTHORIZATION=bearer(user))
        errors = self.post([[other.dtm_id, 1, U1, M1, 1, ""]]).json()["data"]["errors"]
        self.assertEqual(errors[0]["column"], "dtm_id")

    def test_template_download(self):
        response = self.client.get(reverse("computer-import-template"))
        self.assertEqual(response.status_code, 200)
        header = [c.value for c in load_workbook(io.BytesIO(response.content)).active[1]]
        # Shablonda sarlavha QAVSSIZ; qavsli sarlavha (HEADER) esa
        # o'qishda baribir qabul qilinadi - yuqoridagi testlar shu bilan.
        self.assertEqual(
            header, ["dtm_id", "zone_number", "machine_uuid", "mac_address", "number", "inventory_code"]
        )

    def test_view_only_user_cannot_import(self):
        user = factories.make_user(permissions=["devices.view"])
        self.client.credentials(HTTP_AUTHORIZATION=bearer(user))
        self.assertEqual(self.post([[1, 1, U1, "", 1, ""]]).status_code, 403)


SHEET = "xl/worksheets/sheet1.xml"


def rezip(upload, *, sheet=None, extra=None) -> SimpleUploadedFile:
    """Haqiqiy faylning ichini almashtiradi: `sheet(xml) -> xml`, `extra` - qo'shimcha fayllar."""
    source = zipfile.ZipFile(io.BytesIO(upload.read()))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            body = source.read(item)
            if item.filename == SHEET and sheet is not None:
                body = sheet(body.decode()).encode()
            target.writestr(item, body)
        for name, body in (extra or {}).items():
            target.writestr(name, body)
    return SimpleUploadedFile("k.xlsx", buffer.getvalue())


class ComputerImportHardeningTests(TestCase):
    """
    Zararli fayllar (`_check_archive`, `_scan_sheet`): 5 MB lik zip ichidagi
    XML server xotirasini/protsessorini tugatmasligi va 500 bermasligi kerak.
    """

    def setUp(self):
        self.region = factories.make_region()
        factories.make_zone(region=self.region, number=3)
        self.user = factories.make_user(permissions=["devices.view", "devices.manage"])
        self.client = APIClient()
        self.client.credentials(HTTP_AUTHORIZATION=bearer(self.user))
        self.url = reverse("computer-import-excel")
        self.row = [self.region.dtm_id, 3, U1, M1, 1, ""]

    def post(self, upload):
        return self.client.post(self.url, {"file": upload, "dry_run": "true"}, format="multipart")

    def assert_rejected(self, upload, fragment):
        response = self.post(upload)
        self.assertEqual(response.status_code, 400, response.content)
        self.assertIn(fragment, str(response.json()["error"]["details"]["file"]))

    def test_unpacked_size_limit(self):
        # ~21 MB nol - siqilgani bir necha KB.
        bomb = rezip(xlsx([self.row]), extra={"xl/media/bomb.bin": b"\0" * (21 * 1024 * 1024)})
        self.assertLess(bomb.size, 100 * 1024)
        self.assert_rejected(bomb, "juda katta")

    def test_far_row_number_is_rejected_not_dropped(self):
        far = '<row r="20000"><c r="A20000" t="inlineStr"><is><t>1</t></is></c></row></sheetData>'
        self.assert_rejected(rezip(xlsx([self.row]), sheet=lambda x: x.replace("</sheetData>", far)), "qator bor")

    def test_too_many_cells_in_row(self):
        wide = '<row r="3">' + "<c/>" * 1000 + "</row></sheetData>"
        self.assert_rejected(rezip(xlsx([self.row]), sheet=lambda x: x.replace("</sheetData>", wide)), "katak")

    def test_too_many_cells_in_sheet(self):
        grid = "".join(f'<row r="{i}">' + "<c/>" * 200 + "</row>" for i in range(3, 600)) + "</sheetData>"
        self.assert_rejected(rezip(xlsx([self.row]), sheet=lambda x: x.replace("</sheetData>", grid)), "katak")

    def test_entity_expansion_is_400_not_500(self):
        def laughs(xml):
            doctype = '<!DOCTYPE w [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            return re.sub(r"(<\?xml[^>]*\?>)", r"\1" + doctype, xml, count=1).replace("<v>1</v>", "<v>&b;</v>")

        self.assert_rejected(rezip(xlsx([self.row]), sheet=laughs), "o'qilmadi")

    def test_not_a_zip(self):
        self.assert_rejected(SimpleUploadedFile("k.xlsx", b"dtm_id;zone_number\n1;2\n"), "o'qilmadi")

    def test_spoofed_dimension_still_reads(self):
        """`<dimension ref="A1:XFD1048576">` - qatorlar shu o'lchamga to'ldirilmaydi."""
        spoof = rezip(xlsx([self.row]), sheet=lambda x: re.sub(
            r'<dimension ref="[^"]*"', '<dimension ref="A1:XFD1048576"', x))
        response = self.post(spoof)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["data"]["to_create"], 1)

    def test_filled_template_passes(self):
        """Shablonda 5001 ta formatlangan qator bor - qator chegarasi uni to'smasligi kerak."""
        from apps.devices.computer_import import build_template

        wb = load_workbook(io.BytesIO(build_template()))
        ws = wb["Kompyuterlar"]
        for column, value in zip("ABCDEF", self.row):
            ws[f"{column}2"] = value
        buffer = io.BytesIO()
        wb.save(buffer)
        response = self.post(SimpleUploadedFile("k.xlsx", buffer.getvalue()))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["data"]["to_create"], 1)
